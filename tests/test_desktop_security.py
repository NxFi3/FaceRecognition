"""Regression tests for confirmed desktop boundary and lifecycle flaws."""
from concurrent.futures import ThreadPoolExecutor
import csv
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import Mock, patch
import zlib
import base64

import requests
import test_cloud_client as fixture
from src.cloud.client import CloudClient, AccessDenied, CloudError
from src.cloud.offline import Checkpoint
from src.cloud.recovery import import_key
from src.spreadsheet_security import spreadsheet_cell, spreadsheet_row


class SessionSecurityTests(unittest.TestCase):
    setUp = fixture.CloudClientTests.setUp
    envelope = fixture.CloudClientTests.envelope
    transport = fixture.CloudClientTests.transport
    login = fixture.CloudClientTests.login

    def restarted(self):
        return CloudClient(self.config, self.client.state_path, store=self.vault, transport=self.transport)

    def test_normal_shutdown_preserves_valid_offline_subscription(self):
        self.login()
        self.client.close()
        restarted = self.restarted()
        self.assertTrue(restarted.allowed())

    def test_logout_then_shutdown_does_not_resurrect_offline_license(self):
        self.login()
        self.client.logout()
        self.client.close()
        restarted = self.restarted()
        self.assertFalse(restarted.allowed())
        self.assertIsNone(restarted.credential)

    def test_failed_online_logout_still_removes_local_admission(self):
        self.login()
        with patch.object(self.client, 'transport', side_effect=requests.ConnectionError('offline')):
            with self.assertRaises(CloudError):
                self.client.logout()
        self.client.close()
        self.assertFalse(self.restarted().allowed())

    def test_logout_failure_to_delete_vault_does_not_auto_sign_in_again(self):
        self.login()
        with patch.object(self.vault, 'delete'):
            self.client.logout()
        self.client.close()
        self.assertFalse(self.restarted().allowed())

    def test_fresh_explicit_online_login_after_logout_remains_supported(self):
        self.login()
        self.client.logout()
        self.client.login('school@example.com', 'password')
        self.assertTrue(self.client.allowed())
        self.assertTrue(self.restarted().allowed())

    def test_logout_discards_renewal_already_waiting_for_server(self):
        self.login()
        entered, release = threading.Event(), threading.Event()
        original = self.client.transport
        def transport(method, url, **kwargs):
            if url.endswith('/lease/'):
                response = original(method, url, **kwargs)
                entered.set()
                self.assertTrue(release.wait(3))
                return response
            return original(method, url, **kwargs)
        self.client.transport = transport
        with ThreadPoolExecutor(max_workers=1) as executor:
            future = executor.submit(self.client.renew)
            try:
                self.assertTrue(entered.wait(2))
                self.client.logout()
            finally:
                release.set()
            with self.assertRaises(AccessDenied):
                future.result(3)
        self.client.close()
        self.assertFalse(self.restarted().allowed())

    def test_logout_discards_login_that_has_received_lease_but_not_saved_token(self):
        self.login()
        entered, release = threading.Event(), threading.Event()
        original = self.client.accept_lease
        def accept(*args, **kwargs):
            value = original(*args, **kwargs)
            entered.set()
            self.assertTrue(release.wait(3))
            return value
        with patch.object(self.client, 'accept_lease', side_effect=accept):
            with ThreadPoolExecutor(max_workers=1) as executor:
                future = executor.submit(self.client.login, 'school@example.com', 'password')
                try:
                    self.assertTrue(entered.wait(2))
                    self.client.logout()
                finally:
                    release.set()
                with self.assertRaises(AccessDenied):
                    future.result(3)
        self.client.close()
        self.assertIsNone(self.client.credential)
        self.assertFalse(self.restarted().allowed())


class FileBoundarySecurityTests(unittest.TestCase):
    def test_recovery_key_read_is_bounded_before_json_parsing(self):
        class TrackingStream(io.BytesIO):
            requested = []
            def read(self, size=-1):
                self.requested.append(size)
                return super().read(size)
        stream = TrackingStream(b'x' * 100000)
        with patch.object(Path, 'open', return_value=stream), patch('src.cloud.recovery.json.loads') as loads:
            with self.assertRaises(ValueError):
                import_key('large.frkey', 'not-used')
            loads.assert_not_called()
            # A bounded read prevents a multi-gigabyte file allocation.
            self.assertTrue(stream.closed)
            self.assertEqual(stream.requested, [4097])

    def test_protected_checkpoint_rejects_trailing_compressed_data(self):
        vault = fixture.MemoryVault()
        checkpoint = Checkpoint(vault, 'test-installation')
        value = {'time': 1, 'tick': 1, 'wall': 1, 'generation': 1}
        packed = zlib.compress(json.dumps(value).encode()) + b'unexpected-trailing-data'
        vault.set(checkpoint.key, 'z2:' + base64.b64encode(packed).decode())
        with self.assertRaises(ValueError):
            checkpoint.load()

    def test_spreadsheet_formulas_are_text_without_changing_numeric_values(self):
        for value in ('=1+1', '+cmd', '-1+2', '@SUM(A1)', '\t=1+1', '\r=1+1', '\n=1+1', '  =1+1'):
            with self.subTest(value=value):
                self.assertEqual(spreadsheet_cell(value), "'" + value)
        self.assertEqual(spreadsheet_row(['دانش‌آموز', '', 3, -5, 1.25, None]), ['دانش‌آموز', '', 3, -5, 1.25, None])

    def test_gui_report_csv_and_xlsx_escape_formulas_without_changing_counts(self):
        from src import app
        from openpyxl import load_workbook
        window = Mock()
        window.report_start.date.return_value.toString.return_value = '2026-01-01'
        window.report_end.date.return_value.toString.return_value = '2026-01-31'
        rows = [(7, '=CODE', '=NAME', 'پایه', '=MAJOR', '=CLASS', 1, 0, 0, 0, 100.0)]
        daily = [('2026-01-01', 1, 0, 0, 0, 100.0)]
        classes = [(1, '=CLASS', '=MAJOR', 1, 1, 0, 0, 0, 100.0)]
        with tempfile.TemporaryDirectory() as directory:
            for fmt in ('csv', 'xlsx'):
                path = Path(directory) / ('report.' + fmt)
                with patch.object(app, 'get_report_rows', return_value=rows), \
                     patch.object(app, 'get_daily_report_summary', return_value=daily), \
                     patch.object(app, 'get_class_report_summary', return_value=classes), \
                     patch.object(app.QFileDialog, 'getSaveFileName', return_value=(str(path), '')), \
                     patch.object(app.QMessageBox, 'information'), \
                     patch.object(app.QMessageBox, 'critical') as critical:
                    app.MainWindow.export_report(window, fmt)
                    critical.assert_not_called()
                if fmt == 'csv':
                    with path.open(encoding='utf-8-sig', newline='') as stream:
                        output = list(csv.reader(stream))
                    self.assertEqual(output[1][0:2], ["'=CODE", "'=NAME"])
                    self.assertEqual(output[1][5], '1')
                else:
                    workbook = load_workbook(path)
                    self.assertEqual(workbook[app.tr('Students')]['A2'].value, "'=CODE")
                    self.assertEqual(workbook[app.tr('Students')]['B2'].data_type, 's')
                    self.assertEqual(workbook[app.tr('Students')]['F2'].value, 1)
                    self.assertEqual(workbook[app.tr('Classes')]['A2'].value, "'=CLASS")
                    workbook.close()

    def test_gui_event_csv_escapes_synthetic_note_formula(self):
        from src import app
        window = Mock()
        window.report_start.date.return_value.toString.return_value = '2026-01-01'
        window.report_end.date.return_value.toString.return_value = '2026-01-31'
        window.report_events_table.rowCount.return_value = 1
        window.report_events_table.columnCount.return_value = 8
        def item(row, column):
            result = Mock()
            result.data.return_value = 77
            result.text.return_value = '=HYPERLINK("https://invalid")' if column == 6 else 'Synthetic'
            return result
        window.report_events_table.item.side_effect = item
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'events.csv'
            with patch.object(app.QFileDialog, 'getSaveFileName', return_value=(str(path), '')), \
                 patch.object(app.QMessageBox, 'information'), \
                 patch.object(app.QMessageBox, 'critical') as critical:
                app.MainWindow.export_report_events(window)
                critical.assert_not_called()
            with path.open(encoding='utf-8-sig', newline='') as stream:
                values = list(csv.reader(stream))
            self.assertEqual(values[1][0], '77')
            self.assertEqual(values[1][7], "'=HYPERLINK(\"https://invalid\")")

    def test_xlsx_does_not_store_user_input_as_formula(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'report.xlsx'
            workbook = Workbook()
            workbook.active.append(spreadsheet_row(['=HYPERLINK("https://invalid")', 3, 'نام']))
            workbook.save(path)
            restored = load_workbook(path)
            self.assertEqual(restored.active['A1'].data_type, 's')
            self.assertEqual(restored.active['A1'].value, "'=HYPERLINK(\"https://invalid\")")
            self.assertEqual(restored.active['B1'].value, 3)
            restored.close()

    def test_attendance_notification_contains_existing_student_display_fields(self):
        from src.engine.smart_camera_worker import SmartCameraWorker
        worker = SmartCameraWorker()
        payloads = []
        worker.attendance_added.connect(payloads.append)
        with patch('src.engine.smart_camera_worker.mark_attendance', return_value=True) as mark, \
             patch('src.engine.smart_camera_worker.get_student', return_value=(7, 'SYN-7', 'Synthetic', '2026-01-01')):
            worker._mark_attendance_event(7, 'camera_0')
        mark.assert_called_once_with(7, 'camera_0', status='present', notes='')
        self.assertEqual(len(payloads), 1)
        self.assertEqual(payloads[0]['student_code'], 'SYN-7')
        self.assertEqual(payloads[0]['name'], 'Synthetic')
